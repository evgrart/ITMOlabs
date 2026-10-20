"""Reproduce all tables, plots and generated data for variant 11."""
import argparse
import csv
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from scipy import stats

from model import (SIZES, CONFIDENCE, read_variant, characteristics,
                   relative_deviation, autocorrelation, fit_h2,
                   generate_h2, h2_pdf, h2_cdf)

ROOT = Path(__file__).resolve().parent
KEYS = ("mean", "eps90", "eps95", "eps99", "variance", "sigma", "cv")


def prefix_table(x):
    return {n: characteristics(x[:n]) for n in SIZES}


def relative_table(table, references):
    return {n: {k: relative_deviation(row[k], references[n][k]) for k in KEYS}
            for n, row in table.items()}


def save_figures(x, y, p, out):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11,
                         "axes.grid": True, "grid.alpha": 0.22,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "pdf.fonttype": 42})
    blue, orange = "#285C87", "#BC6128"

    def save(fig, name):
        fig.tight_layout()
        fig.savefig(out / f"{name}.pdf", bbox_inches="tight")
        plt.close(fig)

    idx = np.arange(1, len(x) + 1)
    fig, ax = plt.subplots(figsize=(9, 3.8))
    ax.plot(idx, x, color=blue, linewidth=0.85)
    ax.axhline(x.mean(), color=orange, linestyle="--", label="Выборочное среднее")
    ax.set(xlabel="Номер измерения i", ylabel="Значение xᵢ", xlim=(1, 300))
    ax.legend()
    save(fig, "01_source")

    # Sturges: ceil(1 + log2(300)) = 10 equal-width bins.
    bins = int(np.ceil(1 + np.log2(len(x))))
    edges = np.linspace(0, max(x), bins + 1)
    fig, ax = plt.subplots(figsize=(9, 3.8))
    counts, _, bars = ax.hist(x, bins=edges, color=blue, alpha=0.85,
                              edgecolor="white")
    ax.bar_label(bars, fmt="%.0f", fontsize=10)
    ax.set(xlabel="Значение x", ylabel="Частота (число наблюдений)",
           ylim=(0, max(counts) * 1.15))
    save(fig, "02_histogram")

    grid = np.linspace(0, max(x.max(), y.max()) * 1.03, 1600)
    fig, ax = plt.subplots(figsize=(9, 3.8))
    ax.hist(x, bins=edges, density=True, alpha=0.5, color=blue,
            edgecolor="white", label="Исходная ЧП: плотность гистограммы")
    ax.plot(grid, h2_pdf(grid, p), color=orange, linewidth=2, label="Плотность H₂")
    # Average model density in each bin is directly comparable to a histogram.
    avg = np.diff(h2_cdf(edges, p)) / np.diff(edges)
    ax.stairs(avg, edges, color="#383838", linestyle="--",
              label="Средняя плотность H₂ по интервалам")
    ax.set(xlabel="Значение x", ylabel="Плотность", xlim=(0, max(x)))
    ax.legend(fontsize=9)
    save(fig, "03_fit")

    detail_edges = np.arange(0, np.ceil(max(x) / 2) * 2 + 2, 2)
    fig, ax = plt.subplots(figsize=(9, 3.4))
    ax.hist(x, bins=detail_edges, density=True, alpha=0.5, color=blue,
            edgecolor="white", label="Исходная ЧП, ширина интервала 2")
    ax.plot(grid, h2_pdf(grid, p), color=orange, linewidth=2, label="Плотность H₂")
    ax.set(xlabel="Значение x", ylabel="Плотность", xlim=(0, 30))
    ax.legend(fontsize=9)
    save(fig, "08_detail")

    fig, axes = plt.subplots(2, 1, figsize=(9, 5.3), sharex=True, sharey=True)
    for ax, values, name, color in zip(axes, (x, y),
            ("Исходная ЧП", "Сгенерированная ЧП"), (blue, orange)):
        ax.plot(idx, values, linewidth=0.8, color=color, label=name)
        ax.axhline(values.mean(), color=color, linestyle="--", linewidth=1)
        ax.set(ylabel="Значение", xlim=(1, 300))
        ax.legend(loc="upper right")
    axes[1].set_xlabel("Номер значения i")
    save(fig, "04_sequences")

    common_edges = np.linspace(0, max(x.max(), y.max()), bins + 1)
    fig, ax = plt.subplots(figsize=(9, 3.8))
    ax.hist([x, y], bins=common_edges, color=[blue, orange],
            label=["Исходная ЧП", "Сгенерированная ЧП"])
    ax.set(xlabel="Значение", ylabel="Частота (число наблюдений)")
    ax.legend()
    save(fig, "05_histograms")

    fig, ax = plt.subplots(figsize=(9, 3.8))
    lags = np.arange(1, 11)
    ax.bar(lags - 0.18, autocorrelation(x), width=0.36, color=blue,
           label="Исходная ЧП")
    ax.bar(lags + 0.18, autocorrelation(y), width=0.36, color=orange,
           label="Сгенерированная ЧП")
    bound = 1.96 / np.sqrt(len(x))
    ax.axhline(bound, color="#555555", linestyle="--", linewidth=1,
               label="Приближённые границы 95 %")
    ax.axhline(-bound, color="#555555", linestyle="--", linewidth=1)
    ax.axhline(0, color="black", linewidth=0.7)
    ax.set(xlabel="Сдвиг k", ylabel="Коэффициент автокорреляции", xticks=lags)
    ax.legend(fontsize=9, ncol=2)
    save(fig, "06_acf")

    fig, ax = plt.subplots(figsize=(9, 3.8))
    for values, color, label in ((x, blue, "Исходная ЧП"),
                                  (y, orange, "Сгенерированная ЧП")):
        ax.step(np.r_[0, np.sort(values)], np.arange(len(values) + 1) / len(values),
                where="post", color=color, label=label)
    ax.plot(grid, h2_cdf(grid, p), color="#383838", linestyle="--", label="H₂")
    ax.set(xlabel="Значение", ylabel="Функция распределения", xlim=(0, grid[-1]))
    ax.legend()
    save(fig, "07_cdf")
    return edges


def cdf_distance(x, parameters):
    theoretical = h2_cdf(np.sort(x), parameters)
    n = len(x)
    return float(max(np.max(np.arange(1, n + 1) / n - theoretical),
                     np.max(theoretical - np.arange(n) / n)))


def exports(out, x, y, result, edges):
    workbook = Workbook()
    workbook.remove(workbook.active)
    tables = {"data": (["i", "source", "generated"],
                        [[i, a, b] for i, (a, b) in enumerate(zip(x, y), 1)])}
    for kind in ("source", "generated"):
        headers = ["n"] + [v for k in KEYS for v in (k, k + "_deviation_percent")]
        rows = [[n] + [v for k in KEYS for v in
                       (result[kind][n][k], result[kind + "_relative"][n][k])]
                for n in SIZES]
        tables[kind] = (headers, rows)
    tables["autocorrelation"] = (
        ["lag", "source", "generated", "deviation_percent"],
        [[k, result["source_acf"][k - 1], result["generated_acf"][k - 1],
          result["acf_relative"][k - 1]] for k in range(1, 11)])
    tables["confidence_intervals"] = (
        ["sequence", "n", "probability", "half_width", "lower", "upper"],
        [[kind, n, prob, result[kind][n][f"eps{round(prob * 100)}"],
          result[kind][n]["mean"] - result[kind][n][f"eps{round(prob * 100)}"],
          result[kind][n]["mean"] + result[kind][n][f"eps{round(prob * 100)}"]]
         for kind in ("source", "generated") for n in SIZES for prob in CONFIDENCE])
    counts, _ = np.histogram(x, bins=edges)
    probabilities = np.diff(h2_cdf(edges, result["parameters"]))
    tables["histogram"] = (["lower", "upper", "observed_count", "expected_count"],
        [[edges[i], edges[i + 1], int(counts[i]), len(x) * probabilities[i]]
         for i in range(len(counts))]
        + [[edges[-1], "infinity", 0, len(x) * (1 - h2_cdf(edges[-1], result["parameters"]))]])
    tables["parameters"] = (["parameter", "value"],
        list(result["parameters"].items()) + [["seed", result["seed"]],
            ["correlation", result["correlation"]],
            ["source_cdf_distance", result["source_cdf_distance"]]])
    for name, (headers, rows) in tables.items():
        with (out / f"{name}.csv").open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(headers)
            writer.writerows(rows)
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append([v.item() if isinstance(v, np.generic) else v for v in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="285C87")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 32
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(
                30, max(15, len(str(col[0].value)) + 2))
            for cell in col[1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000"
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.save(out / "results_variant11.xlsx")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=ROOT / "!_УИР1_Варианты_с.xlsx")
    parser.add_argument("--seed", type=int, default=11)
    parser.add_argument("--output", type=Path, default=ROOT / "results")
    args = parser.parse_args()
    out = args.output.resolve()
    out.mkdir(parents=True, exist_ok=True)
    x, sheet, coordinate = read_variant(args.input)
    p = fit_h2(x)
    y = generate_h2(len(x), p, args.seed)
    source, generated = prefix_table(x), prefix_table(y)
    source_acf, generated_acf = autocorrelation(x), autocorrelation(y)
    trend = stats.linregress(np.arange(1, 301), x)
    correlation = stats.pearsonr(x, y)
    r = dict(variant=11, sheet=sheet, source_range=coordinate, seed=args.seed,
             parameters=p, source=source, generated=generated,
             source_relative=relative_table(source, {n: source[300] for n in SIZES}),
             generated_relative=relative_table(generated, source),
             generated_internal_relative=relative_table(generated, {n: generated[300] for n in SIZES}),
             source_acf=source_acf.tolist(), generated_acf=generated_acf.tolist(),
             acf_relative=[relative_deviation(a, b) for a, b in zip(generated_acf, source_acf)],
             correlation=float(correlation.statistic),
             trend_slope=float(trend.slope), trend_p=float(trend.pvalue),
             acf_boundary=float(1.96 / np.sqrt(len(x))),
             source_cdf_distance=cdf_distance(x, p),
             generated_cdf_distance=cdf_distance(y, p),
             source_min=float(x.min()), source_max=float(x.max()),
             generated_max=float(y.max()),
             source_above100=int(sum(x > 100)), generated_above100=int(sum(y > 100)),
             model_above100=float(1 - h2_cdf(100, p)))
    (out / "results.json").write_text(json.dumps(r, ensure_ascii=False, indent=2), encoding="utf-8")
    edges = save_figures(x, y, p, out)
    exports(out, x, y, r, edges)
    print(json.dumps({key: value for key, value in r.items() if key not in (
        "source", "generated", "source_relative", "generated_relative", "generated_internal_relative")},
        ensure_ascii=False, indent=2))
    print("Source n=300:", source[300])
    print("Generated n=300:", generated[300])
    print("Relative errors n=300:", r["generated_relative"][300])


if __name__ == "__main__":
    main()
